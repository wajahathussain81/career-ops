#!/usr/bin/env python3
"""Export career-ops tracker to applications.xlsx in the live-jobs.xlsx schema."""
import csv
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
HEADERS = ["#", "Relevance", "Level", "Fits you?", "Region", "Company", "Role",
           "Location", "Link", "Apply? (Y/N)", "Notes"]

# url -> location from scan history
loc_by_url = {}
with open(ROOT / "data/scan-history.tsv", newline="") as f:
    for rec in csv.DictReader(f, delimiter="\t"):
        if rec.get("url") and rec.get("location"):
            loc_by_url[rec["url"].strip()] = rec["location"].strip()

def report_url(report_rel):
    p = ROOT / report_rel
    if not p.exists():
        return ""
    for line in p.read_text().splitlines():
        m = re.match(r"\*\*URL:\*\*\s*(\S+)", line)
        if m:
            return m.group(1)
    return ""

def level_from_role(role):
    r = role.lower()
    if re.search(r"\blead\b|\bstaff\b|\bprincipal\b", r):
        return "Lead+"
    if re.search(r"\bsenior\b|\bsr\.?\b", r):
        return "Senior"
    if re.search(r"\bjunior\b|\bjr\.?\b|new grad|intern\b|\bi\b$", r):
        return "Junior"
    return "Mid"

def fits(score):
    if score is None:
        return "?"
    if score >= 3.5:
        return "Yes"
    if score >= 3.0:
        return "Maybe"
    return "No"

def region_from_location(loc):
    if not loc:
        return "🍁 Canada"  # pipeline is Canada-filtered by config
    canada = ("canada", ", on", ", ab", ", bc", ", qc", "ontario", "alberta",
              "british columbia", "quebec", "remote - canada", "toronto",
              "calgary", "vancouver", "montreal", "waterloo", "oakville")
    l = loc.lower()
    if any(k in l for k in canada):
        return "🍁 Canada"
    if "remote" in l:
        return "Remote"
    return loc.split(",")[-1].strip().title() or "?"

rows = []
for line in (ROOT / "data/applications.md").read_text().splitlines():
    if not line.startswith("|"):
        continue
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    if len(cells) < 10 or cells[0] in ("#", "---") or set(cells[0]) <= {"-"}:
        continue
    num, date, company, via, role, score_s, status, pdf, report, notes = cells[:10]
    m = re.match(r"(\d+(?:\.\d+)?)/5", score_s)
    score = float(m.group(1)) if m else None
    rm = re.search(r"\(((?:\.\./)?reports/[^)]+)\)", report)
    rel = rm.group(1).replace("../", "") if rm else ""
    url = report_url(rel) if rel else ""
    loc = loc_by_url.get(url, "")
    display_company = company if company != "?" else f"Confidential (via {via})"
    note_out = notes
    if status not in ("Evaluated", ""):
        note_out = f"[{status}] {notes}".strip()
    rows.append({
        "num": int(num), "score": score, "level": level_from_role(role),
        "fits": fits(score), "region": region_from_location(loc),
        "company": display_company, "role": role, "loc": loc, "url": url,
        "apply": "Y" if (score or 0) >= 3.5 else "N", "notes": note_out,
    })

rows.sort(key=lambda r: (-(r["score"] or 0), r["num"]))

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Applications"
ws.append(HEADERS)
for c in ws[1]:
    c.font = Font(bold=True)
for r in rows:
    ws.append([r["num"], r["score"], r["level"], r["fits"], r["region"],
               r["company"], r["role"], r["loc"], "", r["apply"], r["notes"]])
    link_cell = ws.cell(row=ws.max_row, column=9)
    if r["url"]:
        link_cell.value = "Open ↗"
        link_cell.hyperlink = r["url"]
        link_cell.font = Font(color="0563C1", underline="single")
    else:
        link_cell.value = ""
widths = [5, 10, 8, 9, 11, 22, 42, 28, 9, 13, 60]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"

out = sys.argv[1]
wb.save(out)
missing_url = sum(1 for r in rows if not r["url"])
missing_loc = sum(1 for r in rows if not r["loc"])
print(f"rows={len(rows)} missing_url={missing_url} missing_location={missing_loc} -> {out}")
